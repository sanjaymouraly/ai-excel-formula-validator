import io
import zipfile

import pytest
from openpyxl import Workbook

from app.core.workbooks import validate_upload
from app.engines.libreoffice import EngineError, LibreOfficeEngine


def xlsx(wb):
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_valid_upload_and_formula():
    wb = Workbook()
    wb.active["A1"] = 10
    wb.active["B1"] = "=SUM(A1:A5)"
    assert validate_upload(xlsx(wb)).active["A1"].value == 10


@pytest.mark.parametrize(
    "formula", ['=WEBSERVICE("https://evil.example")', "=[other.xlsx]Sheet!A1", "=INDIRECT(A1)"]
)
def test_block_unsafe_upload_formulas(formula):
    wb = Workbook()
    wb.active["A1"] = formula
    with pytest.raises(ValueError):
        validate_upload(xlsx(wb))


def test_block_external_relationship_and_zip_paths():
    for name, data in [
        ("../../evil", b"x"),
        ("xl/vbaProject.bin", b"x"),
        ("xl/_rels/workbook.xml.rels", b'<Relationship TargetMode="External" />'),
    ]:
        wb = Workbook()
        source = zipfile.ZipFile(io.BytesIO(xlsx(wb)))
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as out:
            for item in source.infolist():
                if item.filename != name:
                    out.writestr(item.filename, source.read(item.filename))
            out.writestr(name, data)
        with pytest.raises(ValueError):
            validate_upload(stream.getvalue())


def test_limits():
    with pytest.raises(ValueError):
        validate_upload(b"x" * (5 * 1024 * 1024 + 1))
    with pytest.raises(ValueError):
        validate_upload(b"not a zip")
    wb = Workbook()
    wb.active["ZZ1000"] = 1
    with pytest.raises(ValueError):
        validate_upload(xlsx(wb))


def test_upload_calculation_refused_without_isolation(tmp_path):
    engine = LibreOfficeEngine()
    engine.isolated = False
    with pytest.raises(EngineError, match="isolation unavailable"):
        engine.calculate(tmp_path, [], "Data", "H2", trusted=False)


def test_relationship_single_quotes_are_also_blocked():
    wb = Workbook()
    source = zipfile.ZipFile(io.BytesIO(xlsx(wb)))
    stream = io.BytesIO()
    name = "xl/_rels/workbook.xml.rels"
    with zipfile.ZipFile(stream, "w") as out:
        for item in source.infolist():
            if item.filename != name:
                out.writestr(item.filename, source.read(item.filename))
        out.writestr(
            name,
            b"<Relationships><Relationship TargetMode='External' Target='https://evil.example'/></Relationships>",
        )
    with pytest.raises(ValueError, match="External"):
        validate_upload(stream.getvalue())
