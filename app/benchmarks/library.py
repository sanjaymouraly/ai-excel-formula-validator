"""Versioned benchmark definitions with independent business-rule Python oracles."""

import random
from dataclasses import asdict, dataclass
from datetime import date

from openpyxl import Workbook


@dataclass(frozen=True)
class Benchmark:
    id: str
    title: str
    category: str
    difficulty: str
    description: str
    reference: str
    defective: str
    failure: str
    version: int = 1
    target: str = "H2"

    def public(self):
        return {
            **asdict(self),
            "schema": {
                "Data": "A: group/id/date, B: amount/quantity, C: price/actual, D: auxiliary; F2: fixed parameter",
                "Rates": "B2: tax rate",
            },
            "generator": "Case 0 is the coincidental sample; case 1 is a controlled adversarial mutation; later cases use seed+case and boundary/expanded inputs.",
            "expected_behavior": "Reference must agree with independent Python business-rule oracle before candidate is scored.",
        }


BENCHMARKS = [
    Benchmark(
        "sales_total",
        "Sales • omitted final row",
        "Sales",
        "easy",
        "Total all sales, including newly appended records, in B2:B100.",
        "=SUM(B2:B100)",
        "=SUM(B2:B5)",
        "Last row starts at zero; nonzero final rows and expansion expose the omission.",
    ),
    Benchmark(
        "wrong_column",
        "Revenue • wrong column",
        "Sales",
        "easy",
        "Total revenue from B, not the coincidentally equal values in C.",
        "=SUM(B2:B100)",
        "=SUM(C2:C100)",
        "B and C match only in the baseline.",
    ),
    Benchmark(
        "hardcoded",
        "Revenue • hardcoded result",
        "Sales",
        "easy",
        "Sum revenue in B2:B100 as inputs change.",
        "=SUM(B2:B100)",
        "=100",
        "A constant matches only the original total.",
    ),
    Benchmark(
        "aggregation",
        "KPI • SUM versus AVERAGE",
        "KPI",
        "easy",
        "Compute the arithmetic average of numeric B cells, ignoring blanks.",
        "=AVERAGE(B2:B100)",
        "=SUM(B2:B100)",
        "All-zero baseline hides the aggregation error.",
    ),
    Benchmark(
        "regional",
        "Regional revenue • condition",
        "Conditional aggregation",
        "medium",
        "Sum B only for rows whose region in A is East.",
        '=SUMIF(A2:A100,"East",B2:B100)',
        '=SUMIF(A2:A100,"West",B2:B100)',
        "Balanced sample regions hide an inverted condition.",
    ),
    Benchmark(
        "discount_tax",
        "Discounts & multi-sheet tax",
        "Finance",
        "medium",
        "Sum quantity B times unit price C, apply discount F2 and tax Rates!B2.",
        "=SUMPRODUCT(B2:B100,C2:C100)*(1-$F$2)*(1+Rates!B2)",
        "=SUMPRODUCT(B2:B100,C2:C100)*(1+Rates!B2)",
        "Zero initial discount hides a missing deduction.",
    ),
    Benchmark(
        "profit",
        "Operating profit margin",
        "Finance",
        "medium",
        "Net margin equals (total revenue B minus total expense C) divided by revenue; return 0 when revenue is 0.",
        "=IFERROR((SUM(B2:B100)-SUM(C2:C100))/SUM(B2:B100),0)",
        "=IFERROR(SUM(C2:C100)/SUM(B2:B100),0)",
        "Equal cost and profit in the baseline mask a misleading margin.",
    ),
    Benchmark(
        "inventory",
        "Inventory • reorder count",
        "Inventory",
        "easy",
        "Count items with stock B strictly below 10.",
        '=COUNTIF(B2:B100,"<10")',
        '=COUNTIF(B2:B100,"<=10")',
        "Boundary stock of 10 exposes inclusive/exclusive error.",
    ),
    Benchmark(
        "payroll",
        "Payroll • fixed hourly rate",
        "Payroll",
        "medium",
        "Pay row 3 hours B3 at the fixed rate F2; formula is the copy-down result from row 2.",
        "=B3*$F$2",
        "=B3*F3",
        "F3 happens to equal F2, until the next record changes.",
        target="H3",
    ),
    Benchmark(
        "delivery",
        "Supplier on-time deliveries",
        "Suppliers",
        "medium",
        "Count supplier Acme deliveries taking at most 5 days (B).",
        '=COUNTIFS(A2:A100,"Acme",B2:B100,"<=5")',
        '=COUNTIFS(A2:A100,"Acme",B2:B100,"<5")',
        "Deliveries taking exactly five days expose the cutoff.",
    ),
    Benchmark(
        "invoices",
        "Outstanding customer balances",
        "Invoices",
        "easy",
        "Total invoice amounts B less payments C, including credits.",
        "=SUM(B2:B100)-SUM(C2:C100)",
        "=SUM(B2:B100)",
        "Initially there are no payments.",
    ),
    Benchmark(
        "budget",
        "Budget versus actual variance",
        "Budget",
        "easy",
        "Compute actual C minus budget B.",
        "=SUM(C2:C100)-SUM(B2:B100)",
        "=SUM(B2:B100)-SUM(C2:C100)",
        "A balanced initial budget hides the reversed subtraction.",
    ),
    Benchmark(
        "weighted",
        "Weighted KPI",
        "KPI",
        "difficult",
        "Compute weighted mean of scores B with weights C; return zero if total weight is zero.",
        "=IFERROR(SUMPRODUCT(B2:B100,C2:C100)/SUM(C2:C100),0)",
        "=AVERAGE(B2:B100)",
        "Equal initial weights hide an unweighted average.",
    ),
    Benchmark(
        "lookup",
        "Customer lookup • first match",
        "Lookup",
        "difficult",
        "Look up key F2 in A and return the first matching B, or text missing. Preserve first-match behavior with duplicates and reordered inputs.",
        '=IFERROR(VLOOKUP(F2,A2:B100,2,FALSE),"missing")',
        '=IFERROR(VLOOKUP(F2,A2:B100,2,TRUE),"missing")',
        "Unsorted, duplicate and missing keys expose approximate-match lookup.",
    ),
    Benchmark(
        "index_match",
        "Inventory INDEX/MATCH",
        "Lookup",
        "medium",
        "Return B for first exact key F2 in A, otherwise missing.",
        '=IFERROR(INDEX(B2:B100,MATCH(F2,A2:A100,0)),"missing")',
        '=IFERROR(INDEX(B2:B100,MATCH(F2,A2:A100,1)),"missing")',
        "Approximate matching fails on absent or reordered keys.",
    ),
    Benchmark(
        "dates",
        "Month-end invoice boundary",
        "Dates",
        "difficult",
        "Sum invoices B dated in January 2024, including January 31 and excluding February 1.",
        '=SUMIFS(B2:B100,A2:A100,">="&DATE(2024,1,1),A2:A100,"<"&DATE(2024,2,1))',
        '=SUMIFS(B2:B100,A2:A100,">="&DATE(2024,1,1),A2:A100,"<"&DATE(2024,1,31))',
        "Month-end dates reveal an off-by-one cutoff.",
    ),
    Benchmark(
        "negatives",
        "Net revenue • credits & blanks",
        "Edge cases",
        "medium",
        "Sum all revenue, including negative credits, zero and blank rows.",
        "=SUM(B2:B100)",
        '=SUMIF(B2:B100,">0",B2:B100)',
        "Positive-only aggregation silently discards refunds.",
    ),
    Benchmark(
        "nested_if",
        "Nested KPI classification",
        "Logic",
        "medium",
        "Classify B2: high at 100 or more, medium at 50 or more, otherwise low.",
        '=IF(B2>=100,"high",IF(B2>=50,"medium","low"))',
        '=IF(B2>100,"high",IF(B2>50,"medium","low"))',
        "Exact thresholds reveal incorrect strict comparisons.",
    ),
    Benchmark(
        "errors",
        "Error propagation",
        "Edge cases",
        "medium",
        "Compute B2/C2, preserving the division-by-zero error.",
        "=B2/C2",
        "=IFERROR(B2/C2,0)",
        "A zero denominator must remain #DIV/0!, not a misleading zero.",
    ),
]
LIBRARY = {b.id: b for b in BENCHMARKS}


@dataclass
class Dataset:
    cells: dict[str, object]
    expected: object
    mutation: str
    case: int
    seed: int

    def workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(
            [
                "Region / Key / Date",
                "Revenue / Qty",
                "Cost / Weight",
                "Auxiliary",
                None,
                "Parameter",
                None,
                "Result",
            ]
        )
        rates = wb.create_sheet("Rates")
        rates["A2"], rates["B2"] = "Tax", 0.18
        for address, value in self.cells.items():
            if "!" in address:
                sheet, cell = address.split("!")
                wb[sheet][cell] = value
            else:
                ws[address] = value
        ws.freeze_panes = "A2"
        for col in ("A", "B", "C", "F", "H"):
            ws.column_dimensions[col].width = 24
        return wb


def generate(key: str, seed: int, case: int) -> Dataset:
    rng = random.Random(seed + case * 100003)
    base = case == 0
    n = 5 if base else (8 if case == 1 else rng.randint(3, 20))
    b = [10, 20, 30, 40, 0] if base else [rng.randint(-40, 200) for _ in range(n)]
    a = ["East"] * n
    c = list(b)
    extra = {"F2": 0.0}
    expected = 0
    if key in {"sales_total", "wrong_column", "hardcoded"}:
        if case == 1 and key == "sales_total":
            b = [10, 20, 30, 40, 7]
            a = ["East"] * 5
        c = list(b) if base or key == "sales_total" else [v + 7 for v in b]
        expected = sum(b)
    elif key == "aggregation":
        b = [0] * n if base else [10, 20, 30] + [0] * (n - 3)
        expected = sum(b) / len(b)
    elif key == "regional":
        a = ["East", "West"] * ((n + 1) // 2)
        b = [10, 10, 20, 20, 0] if base else b
        expected = sum(v for k, v in zip(a, b) if k == "East")
    elif key == "discount_tax":
        b = [max(0, v) for v in b]
        c = [rng.randint(1, 40) for _ in b]
        extra["F2"] = 0 if base else 0.15
        expected = sum(v * w for v, w in zip(b, c)) * (1 - extra["F2"]) * 1.18
    elif key == "profit":
        c = [v / 2 for v in b] if base else [v * 0.7 for v in b]
        expected = (sum(b) - sum(c)) / sum(b) if sum(b) else 0
    elif key == "inventory":
        b = [1, 9, 11, 20, 0] if base else [10, 10, 0, 11, 9]
        expected = sum(v < 10 for v in b)
    elif key == "payroll":
        extra.update({"F2": 25, "F3": 25 if base else 99})
        expected = b[1] * 25
    elif key == "delivery":
        a = ["Acme"] * n
        b = [1, 2, 4, 6, 8] if base else [5, 5, 1, 8]
        expected = sum(v <= 5 for v in b)
    elif key == "invoices":
        c = [0] * n if base else [abs(v) // 2 for v in b]
        expected = sum(b) - sum(c)
    elif key == "budget":
        c = list(b) if base else [v + 10 for v in b]
        expected = sum(c) - sum(b)
    elif key == "weighted":
        c = [1] * n if base else [1, 9] + [2] * (n - 2)
        expected = sum(v * w for v, w in zip(b, c)) / sum(c)
    elif key in {"lookup", "index_match"}:
        a = [1, 2, 3, 4, 5] if base else ([5, 1, 3, 3, 9] if case % 3 else [5, 1, 3, 3, 9][::-1])
        b = [10, 20, 30, 40, 50]
        extra["F2"] = 3 if base or case % 2 == 0 else 4
        expected = next((v for k, v in zip(a, b) if k == extra["F2"]), "missing")
    elif key == "dates":
        epoch = date(1899, 12, 30)
        dates = (
            [date(2024, 1, 1), date(2024, 1, 15), date(2023, 12, 31)]
            if base
            else [
                date(2024, 1, 31),
                date(2024, 2, 1),
                date(2024, 1, 1),
                date(2024, 12, 31),
                date(2025, 1, 1),
            ]
        )
        a = [(d - epoch).days for d in dates]
        b = [10] * (len(a))
        expected = sum(v for d, v in zip(dates, b) if date(2024, 1, 1) <= d < date(2024, 2, 1))
    elif key == "negatives":
        b = [10, 20, 0, None, 30] if base else [100, -50, None, 0, -10]
        expected = sum(v for v in b if v is not None)
    elif key == "nested_if":
        b = [75] if base else [[100, 50, 49, 101][(case - 1) % 4]]
        expected = "high" if b[0] >= 100 else "medium" if b[0] >= 50 else "low"
    elif key == "errors":
        b, c = [10], [2 if base else 0]
        expected = 5 if base else "#DIV/0!"
    else:
        raise KeyError(key)
    cells = dict(extra)
    for col, values in (("A", a), ("B", b), ("C", c)):
        cells.update({f"{col}{i + 2}": v for i, v in enumerate(values)})
    return Dataset(
        cells,
        expected,
        "Coincidental original sample"
        if base
        else LIBRARY[key].failure
        + (
            " Controlled boundary mutation."
            if case == 1
            else f" Seeded variant {case}; includes growth/edge cases."
        ),
        case,
        seed,
    )
