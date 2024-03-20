import hashlib
import json
import logging
import shutil
import tempfile
import time
import uuid
from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.library import LIBRARY, generate
from app.core.analysis import analyze
from app.core.comparison import equal
from app.core.config import DATA
from app.core.models import EvaluationRequest
from app.core.workbooks import check_target
from app.db.store import Store
from app.engines.libreoffice import EngineError, LibreOfficeEngine


class Evaluator:
    def __init__(self, store=None, engine=None):
        self.store = store or Store()
        self.engine = engine or LibreOfficeEngine()

    def evaluate(self, request: EvaluationRequest):
        start = time.monotonic()
        task = LIBRARY.get(request.benchmark)
        upload = None
        if request.upload_id:
            if not request.upload_id.isalnum() or len(request.upload_id) != 32:
                raise ValueError("Invalid upload identifier")
            upload = DATA / "uploads" / f"{request.upload_id}.xlsx"
            if not upload.exists():
                raise ValueError("Upload not found")
            if not request.reference:
                raise ValueError("Uploaded workbooks require a reference formula")
        elif not task:
            raise ValueError("Unknown benchmark")
        reference = request.reference or task.reference
        sheet = request.sheet if upload else "Data"
        target = request.target if upload else task.target
        sheets = load_workbook(upload).sheetnames if upload else ["Data", "Rates"]
        inspection = analyze(request.candidate, sheets)
        ref_inspection = analyze(reference, sheets)
        run_id = uuid.uuid4().hex
        artifact = DATA / "runs" / run_id
        artifact.mkdir(parents=True)
        result = {
            "id": run_id,
            "request": request.model_dump(),
            "reference": reference,
            "benchmark_version": task.version if task else "uploaded-unverified",
            "inspection": inspection,
            "backend": self.engine.info(),
            "status": "Inconclusive",
            "passed": 0,
            "failed": 0,
            "reference_failures": 0,
            "cases": [],
            "counterexample": None,
            "warnings": [
                "Passing finite tests is not proof of formula equivalence.",
                "LibreOffice is not Microsoft Excel; see supported-function documentation.",
            ],
            "artifacts": [],
        }
        if upload:
            result["warnings"].append(
                "Custom reference has NO independent business-rule oracle; correctness is relative to the supplied reference only. Mutations affect explicitly selected numeric cells, not row growth."
            )
        if ref_inspection["errors"] or ref_inspection["unsupported"]:
            result["status"] = "Reference Failure"
            result["error"] = "Reference rejected by static checks"
        elif inspection["errors"]:
            result["status"] = "Syntax Error"
        elif inspection["unsupported"]:
            result["status"] = "Unsupported Function"
        else:
            try:
                self._execute(request, task, upload, reference, sheet, target, artifact, result)
            except (EngineError, ValueError, KeyError) as exc:
                result["status"] = "Inconclusive"
                result["error"] = str(exc)
        result["latency_seconds"] = round(time.monotonic() - start, 3)
        result["accuracy"] = (
            result["passed"] / (result["passed"] + result["failed"])
            if result["passed"] + result["failed"]
            else None
        )
        if result["status"] in {"Reference Failure", "Unsupported Function", "Inconclusive"}:
            result["accuracy"] = None
        self.store.save_run(result)
        logging.getLogger("verifier.audit").info(
            json.dumps(
                {
                    "event": "evaluation_completed",
                    "run_id": run_id,
                    "status": result["status"],
                    "passed": result["passed"],
                    "failed": result["failed"],
                    "duration_seconds": result["latency_seconds"],
                }
            )
        )
        return result

    def _execute(self, req, task, upload, ref, sheet, target, artifact, result):
        import random

        with tempfile.TemporaryDirectory(prefix="formula-") as folder:
            work = Path(folder)
            filenames, datasets = [], []
            original = load_workbook(upload) if upload else None
            if original:
                if sheet not in original.sheetnames:
                    raise ValueError("Selected worksheet does not exist")
                check_target(original[sheet], target)
                for cell in req.mutation_cells:
                    check_target(original[sheet], cell)
                    if (
                        cell == target
                        or not isinstance(original[sheet][cell].value, (int, float))
                        or isinstance(original[sheet][cell].value, bool)
                    ):
                        raise ValueError(
                            "Mutation cells must be numeric input cells, not the target"
                        )
                if not req.mutation_cells:
                    raise ValueError(
                        "Select numeric mutation cells for multi-case upload evaluation"
                    )
            for i in range(req.cases):
                dataset = generate(task.id, req.seed, i) if not upload else None
                datasets.append(dataset)
                changes = {}
                for role, formula in (("reference", ref), ("candidate", req.candidate)):
                    wb = load_workbook(upload) if upload else dataset.workbook()
                    ws = wb[sheet]
                    if upload and i:
                        rng = random.Random(req.seed + i * 100003)
                        for cell in req.mutation_cells:
                            old = ws[cell].value
                            value = 0 if i == 1 else rng.randint(-100, 300)
                            ws[cell] = value
                            changes[cell] = {"before": old, "after": value}
                    ws[target] = formula
                    wb.calculation.fullCalcOnLoad = True
                    name = f"case_{i:03}_{role}.xlsx"
                    wb.save(work / name)
                    filenames.append(name)
                if dataset:
                    baseline = generate(task.id, req.seed, 0).cells
                    changes = {
                        k: {"before": baseline.get(k), "after": v}
                        for k, v in dataset.cells.items()
                        if baseline.get(k) != v
                    }
                datasets[i] = {
                    "oracle": dataset.expected if dataset else None,
                    "mutation": dataset.mutation
                    if dataset
                    else "Original upload"
                    if i == 0
                    else "Controlled numeric input mutation",
                    "changes": changes,
                }
            values = self.engine.calculate(work, filenames, sheet, target, trusted=not bool(upload))
            for i in range(req.cases):
                expected, actual = values[2 * i : 2 * i + 2]
                ds = datasets[i]
                reference_ok = upload is not None or equal(
                    expected, ds["oracle"], req.atol, req.rtol
                )
                # LO reports parser errors through cells, not always process exit codes.
                ref_ok = reference_ok and not (
                    isinstance(expected, str) and expected in {"#NAME?", "#REF!"}
                )
                passed = ref_ok and equal(actual, expected, req.atol, req.rtol)
                row = {
                    "index": i,
                    "seed": req.seed,
                    "mutation": ds["mutation"],
                    "expected": expected,
                    "actual": actual,
                    "oracle": ds["oracle"],
                    "reference_validated": bool(not upload and ref_ok),
                    "reference_ok": ref_ok,
                    "passed": passed,
                    "changes": ds["changes"],
                }
                result["cases"].append(row)
                if not ref_ok:
                    result["reference_failures"] += 1
                    result["status"] = "Reference Failure"
                    break
                if passed:
                    result["passed"] += 1
                else:
                    result["failed"] += 1
                    if result["counterexample"] is None:
                        row["fingerprint"] = hashlib.sha256(
                            json.dumps(
                                {
                                    "benchmark": req.benchmark,
                                    "reference": ref,
                                    "candidate": req.candidate,
                                    "expected": expected,
                                    "actual": actual,
                                    "changes": ds["changes"],
                                },
                                sort_keys=True,
                                default=str,
                            ).encode()
                        ).hexdigest()[:16]
                        row["explanation"] = (
                            f"{ds['mutation']} Reference returned {expected!r}; candidate returned {actual!r}. The original sample {'matched' if result['cases'][0]['passed'] else 'already failed'}. Static warnings are separate from this execution evidence."
                        )
                        row["suggested_correction"] = ref
                        result["counterexample"] = dict(row)
                        for role in ("reference", "candidate"):
                            name = f"case_{i:03}_{role}.xlsx"
                            shutil.copy(work / name, artifact / name)
                            result["artifacts"].append(name)
                        if (
                            i
                            and req.minimize
                            and not upload
                            and task.id
                            in {
                                "sales_total",
                                "wrong_column",
                                "hardcoded",
                                "negatives",
                                "aggregation",
                            }
                        ):
                            self._minimize(work, artifact, i, sheet, target, req, result)
                if i == 0:
                    for role in ("reference", "candidate"):
                        name = f"case_000_{role}.xlsx"
                        shutil.copy(work / name, artifact / name)
                        if name not in result["artifacts"]:
                            result["artifacts"].append(name)
            if result["status"] != "Reference Failure":
                # Calc cell errors such as #N/A and #VALUE! can be valid formula
                # outcomes, not proof of malformed syntax. Preserve execution
                # mismatch evidence without guessing a syntax diagnosis.
                result["status"] = (
                    "Counterexample Found" if result["failed"] else "Passed Tested Cases"
                )

    def _minimize(self, work, artifact, index, sheet, target, req, result):
        # Greedy bounded delta reduction: remove numeric input cells while preserving
        # a genuine differential result. Candidate and reference remain identical-input.
        source = load_workbook(work / f"case_{index:03}_reference.xlsx")
        coords = [
            c.coordinate
            for row in source[sheet]
            for c in row
            if c.coordinate != target
            and isinstance(c.value, (int, float))
            and not isinstance(c.value, bool)
        ]
        attempts, accepted = 0, 0
        for chunk in [coords[: len(coords) // 2], coords[len(coords) // 2 :]]:
            if not chunk:
                continue
            trials = []
            for role, formula in (("reference", result["reference"]), ("candidate", req.candidate)):
                trial = (
                    load_workbook(work / f"case_{index:03}_reference.xlsx")
                    if not accepted
                    else load_workbook(artifact / "minimal_reference.xlsx")
                )
                for coord in chunk:
                    trial[sheet][coord] = None
                trial[sheet][target] = formula
                name = f"min_{attempts}_{role}.xlsx"
                trial.save(work / name)
                trials.append(name)
            attempts += 1
            expected, actual = self.engine.calculate(work, trials, sheet, target)
            if (
                not equal(actual, expected, req.atol, req.rtol)
                and expected is not None
                and not (isinstance(expected, str) and expected.startswith("#"))
            ):
                accepted += 1
                for role, name in zip(("reference", "candidate"), trials):
                    shutil.copy(work / name, artifact / f"minimal_{role}.xlsx")
                result["counterexample"]["minimized"] = {
                    "expected": expected,
                    "actual": actual,
                    "method": "bounded greedy cell deletion; reference re-executed, no independent minimized oracle",
                }
        if accepted:
            result["artifacts"] += ["minimal_reference.xlsx", "minimal_candidate.xlsx"]

    def replay(self, run_id):
        result = self.store.get(run_id)
        if not result or not result["counterexample"]:
            raise ValueError("Run has no saved counterexample")
        case = result["counterexample"]
        names = [f"case_{case['index']:03}_{r}.xlsx" for r in ("reference", "candidate")]
        req = EvaluationRequest(**result["request"])
        with tempfile.TemporaryDirectory(prefix="replay-") as directory:
            path = Path(directory)
            for name in names:
                shutil.copy(DATA / "runs" / run_id / name, path / name)
            task = LIBRARY.get(req.benchmark)
            expected, actual = self.engine.calculate(
                path,
                names,
                req.sheet if req.upload_id else "Data",
                req.target if req.upload_id else task.target,
                trusted=not bool(req.upload_id),
            )
        return {
            "run_id": run_id,
            "case": case["index"],
            "seed": req.seed,
            "expected": expected,
            "actual": actual,
            "reproduced": equal(expected, case["expected"], req.atol, req.rtol)
            and equal(actual, case["actual"], req.atol, req.rtol)
            and not equal(actual, expected, req.atol, req.rtol),
        }
