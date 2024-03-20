"""Validate ZIP envelope before openpyxl; never execute upload macros or cached values."""

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from app.core.analysis import analyze
from app.core.config import MAX_CELLS, MAX_UPLOAD


def validate_upload(content: bytes):
    if len(content) > MAX_UPLOAD:
        raise ValueError("XLSX exceeds 5 MiB limit")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 300 or sum(i.file_size for i in infos) > 25 * 1024 * 1024:
                raise ValueError("Workbook ZIP complexity limit exceeded")
            names = [i.filename for i in infos]
            if len(names) != len(set(names)):
                raise ValueError("Duplicate ZIP entries blocked")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("Not an XLSX workbook")
            for name in names:
                low = name.lower()
                if (
                    ".." in Path(name).parts
                    or name.startswith("/")
                    or any(
                        x in low
                        for x in (
                            "vbaproject",
                            "externallinks",
                            "embeddings",
                            "connections",
                            "querytables",
                            "customui",
                        )
                    )
                ):
                    raise ValueError(
                        "Macros, external links, embedded objects or unsafe ZIP paths blocked"
                    )
            for info in infos:
                if info.filename.endswith((".xml", ".rels")):
                    data = archive.read(info)
                    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                        raise ValueError("XML entities/DTDs blocked")
                    try:
                        tree = ET.fromstring(data)
                    except ET.ParseError as exc:
                        raise ValueError("Malformed workbook XML") from exc
                    if info.filename.endswith(".rels") and any(
                        node.attrib.get("TargetMode", "").lower() == "external"
                        for node in tree.iter()
                    ):
                        raise ValueError("External relationships blocked")
                    if info.filename.startswith("xl/worksheets/"):
                        cell_count = sum(
                            1 for node in tree.iter() if node.tag.rsplit("}", 1)[-1] == "c"
                        )
                        if cell_count > MAX_CELLS:
                            raise ValueError("Worksheet XML exceeds cell limit")
        wb = load_workbook(io.BytesIO(content), keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Invalid XLSX archive") from exc
    if len(wb.worksheets) > 20 or sum(s.max_row * s.max_column for s in wb) > MAX_CELLS:
        raise ValueError("Workbook exceeds 20 sheets / 20,000 used cells")
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    if not isinstance(cell.value, str):
                        raise ValueError("Array/data-table formulas are not supported")
                    report = analyze(cell.value, wb.sheetnames)
                    if report["errors"] or report["unsupported"]:
                        raise ValueError(
                            f"Unsafe/unsupported formula at {sheet.title}!{cell.coordinate}"
                        )
    return wb


def check_target(sheet, target):
    if not target or ":" in target or "$" in target:
        raise ValueError("Target must be a single A1 cell")
    try:
        row, col = coordinate_to_tuple(target)
    except (ValueError, KeyError) as exc:
        raise ValueError("Invalid target cell") from exc
    if not (1 <= row <= 1000 and 1 <= col <= 100):
        raise ValueError("Target exceeds supported 1000 rows / 100 columns")
    if max(sheet.max_row, row) * max(sheet.max_column, col) > MAX_CELLS:
        raise ValueError("Workbook too complex")
